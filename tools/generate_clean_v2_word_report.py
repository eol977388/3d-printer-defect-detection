"""Convert the clean_v2 Markdown review report to a Word file with embedded figures."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


FIGURES = [
    ("01_train_class_distribution.png", "图1  train类别分布"),
    ("02_val_class_distribution.png", "图2  val类别分布"),
    ("03_train_object_size_distribution.png", "图3  train目标框大小分布"),
    ("04_val_object_size_distribution.png", "图4  val目标框大小分布"),
    ("05_train_bbox_shape_distribution.png", "图5  train目标框宽高及比例分布"),
    ("06_val_bbox_shape_distribution.png", "图6  val目标框宽高及比例分布"),
    ("07_train_rgb_histogram.png", "图7  train RGB直方图"),
    ("08_val_rgb_histogram.png", "图8  val RGB直方图"),
    ("09_brightness_distribution.png", "图9  train与val图像平均亮度分布"),
    ("10_overexposure_ratio_distribution.png", "图10  train与val中过曝像素比例分布"),
    ("11_blur_score_distribution.png", "图11  train与val基于Laplacian方差的模糊度分布"),
]


def set_cell_shading(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), fill)
    props.append(shade)


def set_fonts(document):
    styles = document.styles
    for name, size, bold in (("Normal", 10.5, False), ("Title", 20, True),
                             ("Heading 1", 16, True), ("Heading 2", 13, True)):
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(size)
        style.font.bold = bold
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def add_table(document, lines):
    data = [[part.strip() for part in line.strip().strip("|").split("|")] for line in lines]
    if len(data) > 1 and all(re.fullmatch(r":?-{3,}:?", item) for item in data[1]):
        data.pop(1)
    table = document.add_table(rows=len(data), cols=len(data[0]))
    table.style = "Table Grid"
    for row_index, values in enumerate(data):
        for col_index, value in enumerate(values):
            cell = table.cell(row_index, col_index)
            cell.text = value
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.name = "Microsoft YaHei"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
                    if row_index == 0:
                        run.bold = True
            if row_index == 0:
                set_cell_shading(cell, "D9EAF7")
    document.add_paragraph()


def add_figures(document, figure_dir):
    document.add_heading("分析图表", level=1)
    for filename, caption in FIGURES:
        path = figure_dir / filename
        if not path.is_file():
            raise FileNotFoundError(path)
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(path), width=Cm(16.0))
        cap = document.add_paragraph(caption)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.runs[0].bold = True
        cap.runs[0].font.name = "Microsoft YaHei"
        cap.runs[0]._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--markdown", type=Path, required=True)
    parser.add_argument("--figures", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    lines = args.markdown.read_text(encoding="utf-8").splitlines()
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Cm(2.2)
    section.left_margin = section.right_margin = Cm(2.2)
    set_fonts(document)
    table_buffer, figures_inserted = [], False

    def flush_table():
        nonlocal table_buffer
        if table_buffer:
            add_table(document, table_buffer)
            table_buffer = []

    for line in lines:
        if line.startswith("|"):
            table_buffer.append(line)
            continue
        flush_table()
        stripped = line.strip()
        if stripped.startswith("!["):
            continue
        if stripped == "## 6. 分析图表":
            add_figures(document, args.figures)
            figures_inserted = True
            continue
        if figures_inserted and re.match(r"^\d+\. `figures/", stripped):
            continue
        if not stripped:
            continue
        if stripped.startswith("# "):
            document.add_heading(stripped[2:], 0)
        elif stripped.startswith("## "):
            document.add_heading(re.sub(r"^##\s+", "", stripped), 1)
        elif stripped.startswith("### "):
            document.add_heading(re.sub(r"^###\s+", "", stripped), 2)
        elif stripped.startswith("> "):
            paragraph = document.add_paragraph(stripped[2:])
            paragraph.style = document.styles["Quote"]
        elif stripped.startswith("- "):
            document.add_paragraph(stripped[2:], style="List Bullet")
        elif re.match(r"^\d+\. ", stripped):
            document.add_paragraph(re.sub(r"^\d+\.\s+", "", stripped), style="List Number")
        else:
            document.add_paragraph(stripped.replace("`", ""))
    flush_table()
    if not figures_inserted:
        add_figures(document, args.figures)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.text = "3D打印机缺陷检测项目｜训练前数据分析与复检"
    args.output.parent.mkdir(parents=True, exist_ok=True)
    document.save(args.output)
    print(args.output)


if __name__ == "__main__":
    main()
